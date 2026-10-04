import fitz  # PyMuPDF
import openpyxl
import json
import collections
import re
import sys

def audit_excel(xlsx_path):
    print(f"=== AUDITING EXCEL: {xlsx_path} ===")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb['KKA Kelas X']
    
    classes = []
    current_class = None
    
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 15)]
        val_strs = [str(v).strip() if v is not None else '' for v in vals]
        joined = ' '.join(val_strs)
        
        # Check for class header
        if 'fase e /' in joined.lower():
            class_name = ''
            for v in val_strs:
                m = re.search(r'X\s+(?:TO\s+\d+|TJKT\s+\d+|DKV\s+\d+|RPL)', v)
                if m:
                    class_name = m.group(0).strip()
            
            program = ''
            wali_kelas = ''
            # Scan nearby rows for program & wali kelas
            for offset in range(-2, 4):
                check_r = r + offset
                if 1 <= check_r <= ws.max_row:
                    for cv in [str(ws.cell(check_r, c).value or '').strip() for c in range(1, 15)]:
                        if cv.startswith(':') and 'Teknik' in cv or 'Desain' in cv or 'Rekayasa' in cv:
                            program = cv[1:].strip()
                        elif cv.startswith(':') and any(title in cv for title in ['S.Pd', 'M.Pd', 'S.Sos', 'SE', 'S.Kom']):
                            wali_kelas = cv[1:].strip()
            
            current_class = {
                'class_name': class_name,
                'program': program,
                'wali_kelas': wali_kelas,
                'header_row': r,
                'students': []
            }
            classes.append(current_class)
            continue
            
        c1 = ws.cell(r, 1).value
        c2 = ws.cell(r, 2).value
        if current_class and (isinstance(c1, int) or (isinstance(c1, str) and c1.strip().isdigit())):
            if c2 and isinstance(c2, str) and len(c2.strip()) > 2 and 'NAMA' not in c2.upper():
                student_no = int(c1) if isinstance(c1, int) else int(c1.strip())
                student_name = c2.strip()
                grades = [ws.cell(r, c).value for c in range(3, 12)]
                has_grades = any(g is not None and str(g).strip() != '' for g in grades)
                current_class['students'].append({
                    'row': r,
                    'no': student_no,
                    'name': student_name,
                    'has_grades': has_grades
                })
                
    # Also check Sheet2
    sheet2_students = []
    if 'Sheet2' in wb.sheetnames:
        ws2 = wb['Sheet2']
        for r in range(1, ws2.max_row + 1):
            c1 = ws2.cell(r, 1).value
            c2 = ws2.cell(r, 2).value
            if isinstance(c1, int) or (isinstance(c1, str) and c1.strip().isdigit()):
                if c2 and isinstance(c2, str) and len(c2.strip()) > 2 and 'NAMA' not in c2.upper():
                    sheet2_students.append({
                        'row': r,
                        'no': int(c1) if isinstance(c1, int) else int(c1.strip()),
                        'name': c2.strip()
                    })
                    
    return classes, sheet2_students

def audit_pdf(pdf_path):
    print(f"\n=== AUDITING PDF: {pdf_path} ===")
    doc = fitz.open(pdf_path)
    
    # Page 6: Teacher directory
    page6 = doc[5]
    text6 = page6.get_text('text')
    
    # Parse teachers: each entry has a number, name, subjects, classes, days
    # Let's extract by scanning lines
    lines6 = [l.strip() for l in text6.split('\n') if l.strip()]
    
    # Day schedules: Pages 1 to 5 (Senin, Selasa, Rabu, Kamis, Jumat)
    days = ['SENIN', 'SELASA', 'RABU', 'KAMIS', 'JUMAT']
    class_columns = [
        'X TO 1', 'X TO 2', 'X TO 3', 'X TO 4', 'X TO 5',
        'X TJKT 1', 'X TJKT 2', 'X DKV 1', 'X DKV 2', 'X RPL'
    ]
    
    schedule_by_day = {}
    
    for page_idx in range(5):
        day_name = days[page_idx]
        page = doc[page_idx]
        blocks = page.get_text('blocks')
        schedule_by_day[day_name] = blocks
        
    return text6, schedule_by_day

if __name__ == '__main__':
    xlsx_path = r'C:\Users\vitam\Documents\FORM NILAI ATS GANJIL.xlsx'
    pdf_path = r'C:\Users\vitam\Downloads\JADWAL PELAJARAN YES 2026.2027 (19 AGUSTUS 2026) (Recovered) (Recovered).pdf'
    
    classes, sheet2_students = audit_excel(xlsx_path)
    print(f"Total classes found in Excel: {len(classes)}")
    total_students = sum(len(c['students']) for c in classes)
    print(f"Total students in KKA Kelas X: {total_students}")
    print(f"Total students in Sheet2: {len(sheet2_students)}")
    
    for c in classes:
        print(f"  {c['class_name']}: {len(c['students'])} siswa | Wali: {c['wali_kelas']} | Prog: {c['program']}")
        
    text6, schedule = audit_pdf(pdf_path)
    print("\nPDF Teacher Directory Length:", len(text6))
